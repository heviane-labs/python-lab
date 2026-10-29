import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()

# Sheet 1: Checklist de Devolucao
ws1 = wb.active
ws1.title = "Checklist_Devolucao"
ws1.views.sheetView[0].showGridLines = True

# Colors (Muted Professional Theme - Steel Blue)
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E78")
SECTION_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="1F4E78")
DATA_FONT = Font(name="Calibri", size=11)
BOLD_FONT = Font(name="Calibri", size=11, bold=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# Title Block
ws1.merge_cells("A1:E1")
ws1["A1"] = "CHECKLIST DE DEVOLUÇÃO DE VEÍCULO CORPORATIVO"
ws1["A1"].font = TITLE_FONT
ws1["A1"].alignment = Alignment(vertical="center")

# General Info Block
info_labels = [
    ("A3", "Nome do Colaborador:"), ("C3", "CPF/Matrícula:"),
    ("A4", "Modelo/Marca do Veículo:"), ("C4", "Placa:"),
    ("A5", "Quilometragem (KM):"), ("C5", "Nível de Combustível:"),
    ("A6", "Data da Devolução:"), ("C6", "Responsável pela Recepção:")
]

for pos, text in info_labels:
    ws1[pos] = text
    ws1[pos].font = BOLD_FONT

# Formatting info cells borders/alignments
for row in range(3, 7):
    ws1[f"B{row}"].border = THIN_BORDER
    ws1[f"D{row}"].border = THIN_BORDER

# Header for Checklist Table
headers = ["Item / Componente", "Categoria", "Status de Conservação", "Conforme / Entregue", "Observações / Avarias"]
for col_idx, h in enumerate(headers, 1):
    cell = ws1.cell(row=8, column=col_idx, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

checklist_items = [
    # Documentos e Chaves
    ("Documento do Veículo (CRLV)", "Documentação"),
    ("Chave Principal", "Documentação"),
    ("Chave Reserva", "Documentação"),
    ("Manual do Proprietário e Libreto de Garantia", "Documentação"),
    ("Cartão de Combustível / Sem Parar / Tag", "Documentação"),
    
    # Exterior e Carroceria
    ("Pintura e Lataria (sem amassados ou arranhões)", "Exterior"),
    ("Para-brisas e Vidros (sem trincas/tricados)", "Exterior"),
    ("Faróis, Lanternas e Piscas", "Exterior"),
    ("Espelhos Retrovisores Externos", "Exterior"),
    ("Limpadores de Para-brisa (Palhetas)", "Exterior"),
    ("Estado dos Pneus (Dianteiros e Traseiros)", "Exterior"),
    
    # Interior e Acessórios
    ("Estofamento e Bancos (limpeza/conservação)", "Interior"),
    ("Painel de Instrumentos e Luzes de Alerta", "Interior"),
    ("Ar-Condicionado e Ventilação", "Interior"),
    ("Sistema de Som / Multimídia", "Interior"),
    ("Tapetes Internos", "Interior"),
    
    # Equipamentos Obrigatórios e Segurança
    ("Pneu Estepe (Estado e Calibragem)", "Segurança"),
    ("Macaco e Chave de Roda", "Segurança"),
    ("Triângulo de Sinalização", "Segurança"),
    ("Extintor de Incêndio (se aplicável/validade)", "Segurança"),
    ("Kits de Ferramentas / Primeiros Socorros", "Segurança")
]

start_row = 9
for idx, (item, cat) in enumerate(checklist_items, start=start_row):
    ws1.cell(row=idx, column=1, value=item).font = DATA_FONT
    ws1.cell(row=idx, column=2, value=cat).font = DATA_FONT
    
    # Status default
    c_status = ws1.cell(row=idx, column=3, value="Bom")
    c_status.alignment = Alignment(horizontal="center")
    c_status.font = DATA_FONT
    
    # Conforme default
    c_conf = ws1.cell(row=idx, column=4, value="Sim")
    c_conf.alignment = Alignment(horizontal="center")
    c_conf.font = DATA_FONT
    
    c_obs = ws1.cell(row=idx, column=5, value="")
    c_obs.font = DATA_FONT
    
    for c in range(1, 6):
        ws1.cell(row=idx, column=c).border = THIN_BORDER

end_row = start_row + len(checklist_items) - 1

# Data Validation for Status & Conforme
dv_status = DataValidation(type="list", formula1='"Excelente, Bom, Regular, Danificado, N/A"', allow_blank=True)
dv_conf = DataValidation(type="list", formula1='"Sim, Não, N/A"', allow_blank=True)

ws1.add_data_validation(dv_status)
ws1.add_data_validation(dv_conf)

dv_status.add(f"C{start_row}:C{end_row}")
dv_conf.add(f"D{start_row}:D{end_row}")

# Signatures Block
sig_row = end_row + 3
ws1.cell(row=sig_row, column=1, value="_________________________________________").alignment = Alignment(horizontal="center")
ws1.cell(row=sig_row, column=4, value="_________________________________________").alignment = Alignment(horizontal="center")

ws1.cell(row=sig_row+1, column=1, value="Assinatura do Colaborador").alignment = Alignment(horizontal="center")
ws1.cell(row=sig_row+1, column=1).font = BOLD_FONT

ws1.cell(row=sig_row+1, column=4, value="Assinatura do Responsável (Recebedor)").alignment = Alignment(horizontal="center")
ws1.cell(row=sig_row+1, column=4).font = BOLD_FONT

# Column Widths
widths = {'A': 45, 'B': 18, 'C': 22, 'D': 20, 'E': 40}
for col, width in widths.items():
    ws1.column_dimensions[col].width = width

filename = "Checklist_Entrega_Veiculo_Corporativo.xlsx"
wb.save(filename)
print(f"FILE_CREATED: {filename}")

FILE_CREATED: Checklist_Entrega_Veiculo_Corporativo.xlsx
